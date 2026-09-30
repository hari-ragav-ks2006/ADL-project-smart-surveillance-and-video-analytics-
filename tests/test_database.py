"""Unit tests for Database Persistence, Screenshot capture, and Excel generation."""

from pathlib import Path
import numpy as np
import openpyxl
import pytest
from database import AlertRecord, DatabaseManager
from logger import SurveillanceLogger


def test_database_crud_and_excel(tmp_path):
    """Verify SQLite alert insertion, resolution, screenshot writing, and Excel export."""
    db_file = tmp_path / "test.db"
    screenshots_dir = tmp_path / "screenshots"
    excel_file = tmp_path / "alerts.xlsx"

    db_mgr = DatabaseManager(str(db_file))
    logger_inst = SurveillanceLogger(
        db_manager=db_mgr,
        screenshots_dir=str(screenshots_dir),
        excel_path=str(excel_file),
    )

    # 1. Log an alert with dummy frame
    dummy_frame = np.full((100, 100, 3), 120, dtype=np.uint8)
    record = logger_inst.log_alert(
        alert_type="FACE_COVERED",
        severity="MEDIUM",
        module_name="atm_rules",
        message="Face covered test",
        tracked_ids=[42],
        zone_name="ATM Interaction Zone",
        frame=dummy_frame,
    )

    assert record.id is not None
    assert record.alert_type == "FACE_COVERED"
    assert record.resolution_status == "ACTIVE"
    assert record.screenshot_path is not None
    assert Path(record.screenshot_path).exists()

    # 2. Query alerts
    alerts = logger_inst.query_alerts(severity="MEDIUM")
    assert len(alerts) == 1
    assert alerts[0]["tracked_ids"] == [42]

    # 3. Resolve alert
    resolved = logger_inst.resolve_alert(record.id, resolution_status="RESOLVED")
    assert resolved.resolution_status == "RESOLVED"
    assert resolved.resolved_at is not None

    # 4. Verify Excel file generation and content
    assert excel_file.exists()
    wb = openpyxl.load_workbook(str(excel_file))
    ws = wb.active
    assert ws.title == "Surveillance Alerts"
    assert ws.cell(row=2, column=3).value == "FACE_COVERED"
    assert ws.cell(row=2, column=4).value == "MEDIUM"
    assert ws.cell(row=2, column=9).value == "RESOLVED"
